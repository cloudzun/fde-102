# -*- coding: utf-8 -*-
"""基于《项目汇报PPT模板.pptx》生成西岭项目复盘汇报范例（8 页）"""
import copy
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

TPL = r"D:\vibecoding\dsh\FDE实战课程20260830\项目汇报PPT模板.pptx"
OUT = r"D:\vibecoding\dsh\FDE实战课程20260830\西岭项目复盘汇报-范例-模板版.pptx"

# 模板品牌色
BLUE   = RGBColor(0x00, 0x88, 0xD0)   # 模板主色（标题/表头）
BLUE_D = RGBColor(0x1F, 0x49, 0x7D)   # 主题 dk2，正文强调
BLUE_L = RGBColor(0x11, 0x84, 0xCF)   # 模板辅助蓝
BLUE_P = RGBColor(0x4F, 0x81, 0xBD)   # 主题 accent1，页码/脚注
GREEN  = RGBColor(0x88, 0xDB, 0x29)   # 模板点缀绿
ORANGE = RGBColor(0xE8, 0xA3, 0x3D)   # 警示（未做透）
ORANGE_D = RGBColor(0x5A, 0x3A, 0x05)
RED    = RGBColor(0xC0, 0x50, 0x4D)   # 主题 accent2（风险/缺口）
GREY   = RGBColor(0x7F, 0x7F, 0x7F)
INK    = RGBColor(0x21, 0x2B, 0x36)
LINE   = RGBColor(0xD6, 0xE4, 0xF0)
BG     = RGBColor(0xF2, 0xF8, 0xFD)   # 淡蓝底
BG2    = RGBColor(0xE8, 0xF4, 0xFC)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"

X0, X1 = 1.05, 32.60          # 内容左右边界
W = X1 - X0                    # 31.55
TOP, BOT = 2.75, 16.25         # 内容上下边界（避开左下角徽标 16.69 与右上角徽标 1.9）

prs = Presentation(TPL)
LAYOUT_CONTENT = prs.slide_layouts[1]   # '3_标题和内容'


def set_font(run, size, bold=False, color=INK, font=FONT):
    f = run.font
    f.name = font; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def tb(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = box.text_frame
    tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.1); tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(lines):
        text, size, bold, color = item[0], item[1], item[2], item[3]
        sa = item[4] if len(item) > 4 else 4
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(sa)
        r = p.add_run(); r.text = text; set_font(r, size, bold, color)
    return box


def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if s.has_text_frame:
        s.text_frame.text = ""
    return s


def style_table(tbl, header_rows=1):
    """改为无内置样式的细网格，避免默认蓝色带纹与模板冲突"""
    tblPr = tbl._tbl.tblPr
    tblPr.set('firstRow', '0'); tblPr.set('bandRow', '0')
    for el in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(el)
    el = tblPr.makeelement(qn('a:tableStyleId'), {})
    el.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'  # No Style, Table Grid
    tblPr.append(el)


def page(no, title):
    """内容页：复用模板内容页的标题风格（左对齐 / 垂直居中 / 28pt / #0088D0）"""
    slide = prs.slides.add_slide(LAYOUT_CONTENT)
    tb(slide, 1.05, 0.33, 26.5, 1.94, [(title, 28, True, BLUE)],
       anchor=MSO_ANCHOR.MIDDLE)
    rect(slide, X0, 2.30, 23.5, 0.08, fill=BLUE)
    rect(slide, X0 + 23.5 + 0.05, 2.30, 3.00, 0.08, fill=GREEN)
    tb(slide, 15.0, 17.35, 17.6, 0.7,
       [(f"FDE 政企 AI 落地实战训练营 · 第 3 组　|　第 {no} 页", 11, False, BLUE_P)],
       align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    return slide


# ============================== P1 封面（直接用模板封面，仅换文字） ==============================
s1 = prs.slides[0]
shapes = {sh.name: sh for sh in s1.shapes}

title_shape = shapes['矩形 8']
title_shape.text_frame.paragraphs[0].runs[0].text = "西岭市民服务平台+AI+智能化项目"
for r in title_shape.text_frame.paragraphs[0].runs[1:]:
    r._r.getparent().remove(r._r)

# 副标题（居中，与封面文字同为品牌蓝系）
tb(s1, 7.41, 9.35, 19.05, 1.50, [("复  盘  汇  报", 30, True, BLUE_L)],
   align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
tb(s1, 5.41, 11.05, 23.05, 0.80, [("—— 用 FDE 方法回看一次交付（格式范例）", 14, False, BLUE_P)],
   align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

info = shapes['Text Box 3']
tf = info.text_frame
new_lines = [("第 3 组 　FDE 政企 AI 落地实战训练营", None, None, None),
             ("汇报人：Echo / Delta　组员：（组员姓名）", None, None, None),
             ("2026 年 8 月", None, None, None),
             ("本页为格式范例，项目信息均为虚构教学案例", 12, GREY, False)]
paras = tf.paragraphs
# 复制末段以承载第 4 行，保持原有字体属性
extra = copy.deepcopy(paras[-1]._p)
paras[-1]._p.addnext(extra)
paras = tf.paragraphs
for i, (text, size, color, bold) in enumerate(new_lines):
    p = paras[i]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    if size is not None:
        set_font(p.runs[0], size, bold, color, font=p.runs[0].font.name or FONT)
        p.alignment = PP_ALIGN.CENTER
s1.notes_slide.notes_text_frame.text = (
    "【范例提示】封面沿用模板原设计（青色渐变带 + 品牌蓝标题）。只需替换：项目名、组别、汇报人、日期。"
    "7 分钟汇报不做目录页、不做项目背景铺垫。")

# ============================== P2 结论先行（沿用模板第 2 页，换标题加内容） ==============================
s2 = prs.slides[1]
t2 = {sh.name: sh for sh in s2.shapes}['标题 1']
t2.text_frame.paragraphs[0].runs[0].text = "结论先行：做了什么、解决什么、结论是什么"
for r in t2.text_frame.paragraphs[0].runs[1:]:
    r._r.getparent().remove(r._r)
t2.width = Cm(26.5)   # 收窄，避开模板右上角徽标（x≥28.4）
rect(s2, X0, 2.30, 23.5, 0.08, fill=BLUE)
rect(s2, X0 + 23.55, 2.30, 3.00, 0.08, fill=GREEN)
tb(s2, 15.0, 17.35, 17.6, 0.7,
   [("FDE 政企 AI 落地实战训练营 · 第 3 组　|　第 2 页", 11, False, BLUE_P)],
   align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

rows = [("做了什么", "为市民服务平台交付三项 AI 能力：诉求分类、政策问答、工单分流", BLUE),
        ("解决了什么", "坐席错分返工下降、政策答复可溯源（不再“答错背锅”）、简单件不再占用人工", BLUE),
        ("我们的结论", "Conditional Go（有条件推进）——方向成立，但须先完成本地化部署与人工接管责任落实", BLUE_L)]
y = 2.90
for i, (k, v, c) in enumerate(rows):
    rect(s2, X0, y, W, 2.50, fill=(BG if i % 2 == 0 else WHITE), line=LINE)
    rect(s2, X0, y, 5.40, 2.50, fill=c)
    tb(s2, X0 + 0.25, y + 0.70, 4.90, 1.20, [(k, 16, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
    tb(s2, X0 + 5.75, y + 0.45, W - 6.0, 1.70, [(v, 15, i == 2, INK)], anchor=MSO_ANCHOR.MIDDLE)
    y += 2.70

rect(s2, X0, 11.20, W, 5.00, fill=BG)
rect(s2, X0, 11.20, 0.10, 5.00, fill=RED)
tb(s2, X0 + 0.6, 11.55, W - 1.2, 4.4, [
    ("为什么是“有条件推进”，而不是“直接上”", 15, True, BLUE_D, 8),
    ("① 现有证据只到演示级：真实诉求分布下的准确率、政策覆盖范围均未验证", 14, False, INK, 6),
    ("② 生产条件未落实：本地化部署未做、客户侧人工接管岗位未明确", 14, False, INK, 6),
    ("③ 达成先行条件后，建议转为 Go（见第 7 页）", 14, False, INK, 0)])
s2.notes_slide.notes_text_frame.text = (
    "【范例提示】结论必须放第 2 页，用业务语言（避免 RAG / Agent 等技术词）。"
    "决策四选一：Go / Conditional Go / Continue Pilot / No-Go。")

# ============================== P3 真实问题 ==============================
s = page(3, "真实问题：客户要的 ≠ 真正的业务问题")
colw = (W - 1.0) / 2
rect(s, X0, 2.90, colw, 7.40, fill=BG)
rect(s, X0, 2.90, colw, 0.80, fill=BLUE)
tb(s, X0 + 0.4, 3.00, colw - 0.8, 0.65, [("客户的表面需求（原话）", 14, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
tb(s, X0 + 0.5, 3.95, colw - 1.0, 6.10, [
    ("“我们这个平台每天上万个诉求，坐席分派不过来还老分错；", 14, False, INK, 4),
    ("市民问政策坐席也答不准；简单问题和复杂问题混着办效率低。", 14, False, INK, 4),
    ("你们能不能用 AI 把这一整套都智能化了？”", 14, False, INK, 10),
    ("—— 平台中心 · 陈主任", 12, False, GREY, 0)])
rx = X0 + colw + 1.0
rect(s, rx, 2.90, colw, 7.40, fill=BG)
rect(s, rx, 2.90, colw, 0.80, fill=BLUE_L)
tb(s, rx + 0.4, 3.00, colw - 0.8, 0.65, [("我们重新识别的真实问题", 14, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
tb(s, rx + 0.5, 3.95, colw - 1.0, 6.10, [
    ("不是“上 AI”，而是三个可量化的业务问题：", 14, True, INK, 8),
    ("① 分派不准 → 工单空转、市民重复投诉", 13.5, False, INK, 5),
    ("② 政策答错 → 坐席与平台承担追责风险", 13.5, False, INK, 5),
    ("③ 简单件占人力 → 复杂件得不到足够处理", 13.5, False, INK, 0)])
rect(s, X0, 10.70, W, 5.50, fill=WHITE, line=LINE)
rect(s, X0, 10.70, 0.10, 5.50, fill=BLUE)
tb(s, X0 + 0.6, 11.05, W - 1.2, 4.9, [
    ("复盘发现（当时做得不够的地方）", 15, True, BLUE_D, 8),
    ("方向找对了：把“整套智能化”拆成三个可施工的子场景，是正确的一步。", 14, False, INK, 6),
    ("但当时把三者当成三个并列 demo 交付，忽略了它们的业务闭环关系——", 14, False, INK, 6),
    ("分类的结果本应驱动分流，政策证据本应支撑答复；三者分头做，客户拿到的仍是三件工具，而不是一条流水线。",
     14, False, RED, 0)])
s.notes_slide.notes_text_frame.text = (
    "【范例提示】左栏客户原话（匿名化），右栏你重新识别的问题。当时找错方向也大方写——这是加分项。")

# ============================== P4 八环节复盘 ==============================
s = page(4, "FDE 尺子复盘：八个环节，我们缺在哪一环")
stages = [("office-hours", "确认该做", 0), ("spec", "死磕边界", 1), ("autoplan", "决策拍板", 0),
          ("build", "做出来", 0), ("review", "找隐患", 1), ("qa", "测体验", 0),
          ("ship", "收尾", 1), ("retro", "复盘", 0)]
bw, gap = 3.69, 0.275
for i, (en, cn, st) in enumerate(stages):
    x = X0 + i * (bw + gap)
    color = BLUE if st == 0 else ORANGE
    tc = WHITE if st == 0 else ORANGE_D
    rect(s, x, 2.90, bw, 2.25, fill=color)
    tb(s, x + 0.10, 3.05, bw - 0.20, 0.75, [(en, 12.5, True, tc)], align=PP_ALIGN.CENTER)
    tb(s, x + 0.10, 3.85, bw - 0.20, 0.60, [(cn, 12.5, False, tc)], align=PP_ALIGN.CENTER)
    tb(s, x + 0.10, 4.45, bw - 0.20, 0.60, [("● 做了" if st == 0 else "● 没做透", 11, True, tc)],
       align=PP_ALIGN.CENTER)
    if i < 7:
        tb(s, x + bw, 3.55, gap, 0.8, [("→", 15, True, GREY)], align=PP_ALIGN.CENTER,
           anchor=MSO_ANCHOR.MIDDLE, wrap=False)
rect(s, X0, 5.40, 0.42, 0.42, fill=BLUE)
tb(s, X0 + 0.50, 5.42, 4.0, 0.5, [("做了", 11.5, False, INK)], wrap=False)
rect(s, X0 + 3.40, 5.40, 0.42, 0.42, fill=ORANGE)
tb(s, X0 + 3.90, 5.42, 6.0, 0.5, [("做了没做透", 11.5, False, INK)], wrap=False)
rect(s, X0, 6.25, W, 3.65, fill=BG)
rect(s, X0, 6.25, 0.10, 3.65, fill=ORANGE)
tb(s, X0 + 0.6, 6.55, W - 1.2, 3.25, [
    ("三个“没做透”的环节", 14.5, True, BLUE_D, 5),
    ("spec：写了验收口径，但“数据不出域”只作为约束记录，没有上升为硬 Gate", 13.5, False, INK, 4),
    ("review：查了代码隐患，没有审“生产接管责任”——谁在客户侧接、怎么接", 13.5, False, INK, 4),
    ("ship：交付材料齐全，但没有客户侧运维交接物，客户无法独立运营", 13.5, False, INK, 0)])
rect(s, X0, 10.10, W, 6.15, fill=WHITE, line=RED)
rect(s, X0, 10.10, 0.10, 6.15, fill=RED)
tb(s, X0 + 0.6, 10.40, W - 1.2, 5.6, [
    ("最致命的缺口", 15, True, RED, 10),
    ("spec 阶段没有把“生产化条件”写成硬 Gate。", 15, True, INK, 10),
    ("后果：三个 demo 都停在演示级。验收达标 = 演示可行，但客户无法接管、无法独立运营——", 14, False, INK, 10),
    ("这正是 FDE 与“外包交付”的分水岭：交付物不是“能演示的系统”，而是“客户能接住的能力”。", 14, True, BLUE_D, 12),
    ("如果重来一次：在 spec 就写死——本地化部署完成、人工接管岗位到人、真实脱敏数据复测通过，三项齐备才允许进入 build。",
     14, False, INK, 0)])
s.notes_slide.notes_text_frame.text = (
    "【范例提示】本页是主角，给足时间。八环节横条 + 两色标注，底下写清“最致命缺口 + 后果 + 如果重来一次怎么做”。")

# ============================== P5 证据与假设 ==============================
s = page(5, "证据与假设：哪些已证实，哪些仍是假设")
colw = (W - 1.0) / 2
rect(s, X0, 2.90, colw, 11.90, fill=BG)
rect(s, X0, 2.90, colw, 0.80, fill=BLUE)
tb(s, X0 + 0.4, 3.00, colw - 0.8, 0.65, [("已有证据（可对外）", 15, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
for i, (tag, txt) in enumerate([
        ("已实测", "25 条测试样本上分类准确率达标（演示级）"),
        ("已实测", "政策问答双指标：答案准确率 + 来源可追溯率达标（演示级）"),
        ("已实测", "工单分流三路走通，敏感件零漏判（6 份样本）"),
        ("客户确认", "五类诉求类别清单由客户方确认"),
        ("客户确认", "部门职责库由信息中心提供")]):
    yy = 3.95 + i * 1.90
    c = BLUE if tag == "已实测" else BLUE_L
    tb(s, X0 + 0.4, yy, 2.30, 0.70, [(tag, 11.5, True, c)])
    tb(s, X0 + 2.75, yy - 0.06, colw - 3.2, 1.45, [(txt, 13, False, INK)])
rx = X0 + colw + 1.0
rect(s, rx, 2.90, colw, 11.90, fill=BG)
rect(s, rx, 2.90, colw, 0.80, fill=RED)
tb(s, rx + 0.4, 3.00, colw - 0.8, 0.65, [("仍是假设（不得说成已验证）", 15, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
for i, (tag, txt) in enumerate([
        ("当前假设", "1.2 万件/日真实诉求分布下，准确率仍达标"),
        ("当前假设", "政策库扩至全量后，来源可追溯率不下降"),
        ("待获取", "本地化部署的模型选型、成本与工期"),
        ("待获取", "坐席实际接受度与使用行为数据"),
        ("待获取", "客户侧人工接管的岗位与流程")]):
    yy = 3.95 + i * 1.90
    tb(s, rx + 0.4, yy, 2.90, 0.70, [(tag, 11.5, True, RED)])
    tb(s, rx + 3.35, yy - 0.06, colw - 3.8, 1.45, [(txt, 13, False, INK)])
rect(s, X0, 15.15, W, 1.05, fill=WHITE, line=BLUE)
tb(s, X0 + 0.6, 15.25, W - 1.2, 0.85,
   [("一句话结论：现有证据支持“方向成立”，但不支持“生产可用”；把假设当结论，是复盘最容易翻车的地方。",
     14, True, BLUE_D)], anchor=MSO_ANCHOR.MIDDLE)
s.notes_slide.notes_text_frame.text = (
    "【范例提示】左=已有证据（标状态），右=仍是假设。主动标注假设是加分；把假设说成实测，质询必被打穿。")

# ============================== P6 能力回注 ==============================
s = page(6, "能力回注：这个项目里，什么能沉淀成平台能力")
data = [("客户特定实现（不回注）", "抽象后的通用能力", "层次", "可复用场景"),
        ("市民诉求五分类", "类别可配置的文本分类器", "工具", "银行工单 / 医院投诉"),
        ("西岭 5 份政策问答", "带来源的检索问答契约", "工具", "合规问答 / 制度查询"),
        ("工单三档分流规则", "可配置固定路由工作流", "机制", "各类审批分流"),
        ("敏感件转人工规则", "高风险人工兜底机制", "机制", "风控审批 / 安全事件"),
        ("诉求—政策—部门", "政务业务语义模型", "语义", "其他城市政务热线")]
tx, ty, th = X0, 2.90, 1.18
col_w = [7.60, 8.90, 2.70, 12.35]
tbl = s.shapes.add_table(len(data), 4, Cm(tx), Cm(ty), Cm(W), Cm(th * len(data))).table
for i, w in enumerate(col_w):
    tbl.columns[i].width = Cm(w)
for r in range(len(data)):
    tbl.rows[r].height = Cm(th)
    for c in range(4):
        cell = tbl.cell(r, c)
        cell.text = data[r][c]
        cell.margin_left = cell.margin_right = Cm(0.25)
        cell.margin_top = cell.margin_bottom = Cm(0.06)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE if r == 0 else (WHITE if r % 2 == 1 else BG)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if c == 2 else PP_ALIGN.LEFT
        run = p.runs[0] if p.runs else p.add_run()
        set_font(run, 13.5, r == 0, WHITE if r == 0 else (BLUE_D if c == 0 else INK))
style_table(tbl)
rect(s, X0, 10.40, W, 5.80, fill=WHITE, line=BLUE)
rect(s, X0, 10.40, 0.10, 5.80, fill=GREEN)
tb(s, X0 + 0.6, 10.75, W - 1.2, 5.2, [
    ("回注优先级建议", 14.5, True, BLUE_D, 8),
    ("优先铺“人工兜底机制”与“可配置分类器”——前者三个场景共用、换个客户零改动；后者复用面最广、成本最低。",
     14, False, INK, 8),
    ("业务语义模型价值最高但周期最长，建议作为下一阶段平台专项。", 14, False, INK, 8),
    ("判断标准：换个客户，这东西要不要全改？要改 = 内容（不回注），不用改 = 机制（回注）。", 13.5, False, GREY, 0)])
s.notes_slide.notes_text_frame.text = (
    "【范例提示】每个候选答清：从哪识别 → 抽象成什么 → 哪一层（工具/机制/语义）→ 第二场景 → 谁维护。至少 1 个。")

# ============================== P7 下一步决策 ==============================
s = page(7, "下一步决策：有条件推进（Conditional Go）")
rect(s, X0, 2.90, W, 1.30, fill=BLUE)
tb(s, X0 + 0.6, 3.00, W - 1.2, 1.10,
   [("决　策：Conditional Go —— 三个场景方向成立，可推进；但须先满足下列三项条件，方可转为 Go",
     15.5, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
data2 = [("要素", "内容"),
         ("依据证据", "演示级验收达标（分类 / RAG 双指标 / 路由零漏判）；业务价值方向经客户确认"),
         ("先行条件", "① 完成本地化部署验证　② 明确人工接管岗位到人　③ 用真实脱敏数据复测"),
         ("剩余风险", "真实分布下准确率未知；客户侧运维能力未确认；数据出域方案待信息中心确认"),
         ("补齐责任人", "信息中心 李工（技术接管）/ 平台中心 陈主任（业务验收）/ 承建方（部署实施）"),
         ("下一 Gate", "本地化部署 + 真实数据复测通过后，组织一次评审，转 Go 或调整范围")]
ty2, th2 = 4.50, 1.35
tbl2 = s.shapes.add_table(len(data2), 2, Cm(X0), Cm(ty2), Cm(W), Cm(th2 * len(data2))).table
tbl2.columns[0].width = Cm(5.20); tbl2.columns[1].width = Cm(W - 5.20)
for r in range(len(data2)):
    tbl2.rows[r].height = Cm(th2)
    for c in range(2):
        cell = tbl2.cell(r, c)
        cell.text = data2[r][c]
        cell.margin_left = cell.margin_right = Cm(0.28)
        cell.margin_top = cell.margin_bottom = Cm(0.06)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE if r == 0 else (BG2 if c == 0 else WHITE)
        p = cell.text_frame.paragraphs[0]
        run = p.runs[0] if p.runs else p.add_run()
        set_font(run, 13.5, r == 0 or c == 0, WHITE if r == 0 else (BLUE_D if c == 0 else INK))
style_table(tbl2)
rect(s, X0, 13.00, W, 3.20, fill=BG)
rect(s, X0, 13.00, 0.10, 3.20, fill=RED)
tb(s, X0 + 0.6, 13.30, W - 1.2, 2.7, [
    ("什么证据出现会让我们改判？", 14, True, RED, 6),
    ("→ 真实数据复测准确率显著低于演示水平，或本地化部署成本超出预算，则改判为 Continue Pilot 或 No-Go。",
     13.5, False, INK, 0)])
s.notes_slide.notes_text_frame.text = (
    "【范例提示】决策页必须有四要素：依据证据 / 剩余风险 / 补齐责任人 / 下一 Gate。并补一句“什么证据会改判”。")

# ============================== P8 请求支持 ==============================
s = page(8, "需要决策层支持的事项")
rect(s, X0, 2.90, W, 1.40, fill = BLUE)
tb(s, X0 + 0.6, 3.00, W - 1.2, 1.20,
   [("请求决策层批准以下三项，作为转入下一阶段的前置条件", 15, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
for i, (num, title, desc) in enumerate([
        ("一", "本地化部署资源", "服务器与部署实施人力，用于验证数据不出域的硬约束"),
        ("二", "真实脱敏数据使用授权", "经信息中心与市数据局批准，用于生产分布下的复测"),
        ("三", "人工接管岗位落实", "明确客户侧接管岗位与流程，确保敏感件有人接")]):
    yy = 4.60 + i * 2.95
    rect(s, X0, yy, W, 2.70, fill=BG)
    rect(s, X0, yy, 1.60, 2.70, fill=BLUE_L)
    tb(s, X0, yy + 0.80, 1.60, 1.10, [(num, 16, True, WHITE)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, X0 + 1.95, yy + 0.40, W - 2.3, 1.95, [(title, 15, True, BLUE_D, 4), (desc, 13, False, INK, 0)])
rect(s, X0, 13.60, W, 2.60, fill=WHITE, line=BLUE)
tb(s, X0 + 0.6, 13.80, W - 1.2, 2.2, [
    ("我们的承诺", 14.5, True, BLUE_D, 6),
    ("在获得上述支持后 N 周内，提交本地化部署验证报告与真实数据复测结果；", 14, False, INK, 4),
    ("所有结论仅覆盖取得的证据范围，不夸大交付成果。", 14, False, INK, 0)])
s.notes_slide.notes_text_frame.text = (
    "【范例提示】收尾只讲“要什么支持”，20 秒讲完。不要在此补内容、不要加致谢页。")

prs.save(OUT)
print("已生成:", OUT)
print("页数:", len(prs.slides._sldIdLst))
