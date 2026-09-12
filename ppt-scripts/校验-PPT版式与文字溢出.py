# -*- coding: utf-8 -*-
"""校验模板版范例：越界 / 装饰区 / 文本实际占用范围互相碰撞 / 表格超格"""
import unicodedata, itertools, zipfile, sys
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR

SRC = r"D:\vibecoding\dsh\FDE实战课程20260830\项目汇报PPT模板.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else r"D:\vibecoding\dsh\FDE实战课程20260830\西岭项目复盘汇报-范例-模板版.pptx"
BADGE_TR, BADGE_BL = (28.40, 0.50, 32.50, 1.90), (1.58, 16.69, 4.96, 18.94)

p = Presentation(OUT)
W, H = p.slide_width / 360000, p.slide_height / 360000
print(f"尺寸 {W:.2f}x{H:.2f}cm  页数 {len(p.slides._sldIdLst)}")

# 仅当母版带模板徽标图片（约 x=27.94 / x=1.58）时才检查徽标避让
HAS_BADGE = any(
    sh.left is not None and abs(sh.left / 360000 - 27.94) < 0.3
    for m in p.slide_masters for sh in m.shapes)
print("母版徽标避让检查:", "启用" if HAS_BADGE else "不适用（本版式无模板徽标）")


def wide(ch):
    return unicodedata.east_asian_width(ch) in ("W", "F")


def text_span(sh):
    """估算文本实际占用范围 (l,t,r,b)，返回 None 表示无文字"""
    tf = sh.text_frame
    need, maxw = 0.0, 0.0
    box_w = sh.width / 360000 - 0.25
    paras = [pa for pa in tf.paragraphs if "".join(r.text for r in pa.runs).strip()]
    for idx, para in enumerate(paras):
        txt = "".join(r.text for r in para.runs)
        sz = next((r.font.size.pt for r in para.runs if r.font.size), 14)
        wsum = sum((sz * 0.0353 if wide(c) else sz * 0.0195) for c in txt)
        if tf.word_wrap is False:
            lines = 1
        else:
            lines = 1 if wsum <= box_w else int(wsum // box_w) + 1
        maxw = max(maxw, min(wsum / lines, box_w))
        need += lines * sz * 1.32 / 28.35
        if idx < len(paras) - 1:      # 段末 space_after 是空白，不计入占用
            need += (para.space_after.pt if para.space_after else 0) / 28.35
    if need == 0:
        return None
    y, h = sh.top / 360000, sh.height / 360000
    if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
        t = y + h / 2 - need / 2
    elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
        t = y + h - need
    else:
        t = y + 0.06
    l = sh.left / 360000 + 0.1
    return (l, t, l + maxw, t + need)


total_bad = 0
for i, s in enumerate(p.slides, 1):
    issues, spans = [], []
    for sh in s.shapes:
        if sh.left is None:
            continue
        l, t = sh.left / 360000, sh.top / 360000
        r, b = l + sh.width / 360000, t + sh.height / 360000
        if l < -0.01 or t < -0.01 or r > W + 0.02 or b > H + 0.02:
            issues.append(f"越界 {sh.name} ({l:.2f},{t:.2f})-({r:.2f},{b:.2f})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            sp = text_span(sh)
            if sp:
                spans.append((sh.name, sp))
                if sp[3] > H - 0.1 or sp[2] > W - 0.1 or sp[1] < 0:
                    issues.append(f"文字出界 {sh.name} {tuple(round(v,2) for v in sp)}")
                if i > 1 and HAS_BADGE:
                    for dn, dl, dt, dr, db in (("右上徽标", *BADGE_TR), ("左下徽标", *BADGE_BL)):
                        if sp[0] < dr and sp[2] > dl and sp[1] < db and sp[3] > dt:
                            issues.append(f"文字压{dn} {sh.name} {tuple(round(v,2) for v in sp)}")
    for (n1, s1), (n2, s2) in itertools.combinations(spans, 2):
        ox = min(s1[2], s2[2]) - max(s1[0], s2[0])
        oy = min(s1[3], s2[3]) - max(s1[1], s2[1])
        if ox > 0.15 and oy > 0.05:
            issues.append(f"文字重叠 {n1} ↔ {n2}  ({ox:.2f}×{oy:.2f}cm)")
    tinfo = []
    for sh in s.shapes:
        if not sh.has_table:
            continue
        tb_, cw = sh.table, [c.width / 360000 for c in sh.table.columns]
        cells = []
        for r in range(len(tb_.rows)):
            rh = tb_.rows[r].height / 360000
            for c in range(len(cw)):
                txt = tb_.cell(r, c).text
                if not txt:
                    continue
                sz = 13.5
                wsum = sum((sz * 0.0353 if wide(ch) else sz * 0.0195) for ch in txt)
                avail = cw[c] - 0.55
                if wsum > avail:
                    need = (int(wsum // avail) + 1) * sz * 1.24 / 28.35
                    if need > rh + 0.02:
                        cells.append((r, c, round(need, 2), rh, txt[:24]))
        tinfo.append((len(tb_.rows), round(sum(r.height for r in tb_.rows) / 360000, 2), cells))
        issues += [f"表格超格 r{r}c{c} 需{n}cm 行高{rh} 「{tx}」" for r, c, n, rh, tx in cells]
    total_bad += len(issues)
    print(f"\n[{'OK' if not issues else '!!'}] P{i} 形状{len(s.shapes)} layout={s.slide_layout.name}")
    for x in issues:
        print("    ", x)
    for n, tot, cells in tinfo:
        print(f"     表格 {n}行 高{tot}cm {'无超格' if not cells else ''}")

z0, z1 = zipfile.ZipFile(SRC), zipfile.ZipFile(OUT)
m = lambda z: {n.split('/')[-1] for n in z.namelist() if '/media/' in n and n.split('/')[-1]}
print(f"\n模板媒体保留: {'全部保留' if m(z0) <= m(z1) else '丢失 ' + str(m(z0) - m(z1))}"
      f"  版式数 {sum(1 for n in z1.namelist() if 'slideLayout' in n and n.endswith('.xml'))}")
print(f"合计问题数: {total_bad}")
